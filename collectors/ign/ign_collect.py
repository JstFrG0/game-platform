"""IGN sample collector: saved HTML or interactive browser -> SQLite -> Excel.

The first parser targets the primary game in IGN's __NEXT_DATA__ page object.
It deliberately does not search unrelated recommendations for missing fields.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import sqlite3
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlsplit

ROOT = Path(__file__).resolve().parent


class ScriptParser(HTMLParser):
    """Extract script bodies without running any JavaScript."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=False)
        self.active = None
        self.parts = []
        self.scripts = []

    def handle_starttag(self, tag, attrs):
        if tag == "script":
            self.active = dict(attrs)
            self.parts = []

    def handle_data(self, data):
        if self.active is not None:
            self.parts.append(data)

    def handle_endtag(self, tag):
        if tag == "script" and self.active is not None:
            self.scripts.append((self.active, "".join(self.parts)))
            self.active = None


def ign_url(value: str, kind: str) -> str:
    url = urljoin("https://www.ign.com", value)
    parsed = urlsplit(url)
    if (parsed.scheme != "https" or parsed.hostname not in {"www.ign.com", "ign.com"}
            or not parsed.path.startswith(f"/{kind}/") or parsed.username or parsed.password):
        raise ValueError(f"Expected an IGN /{kind}/ URL: {value}")
    return "https://www.ign.com" + parsed.path.rstrip("/")


def as_date(value):
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10]).isoformat()
    except ValueError:
        return None


def unique(values):
    return list(dict.fromkeys(v for v in values if v))


def resolve_review_object(page, primary, review_url):
    """One manually verified article exception; never infer from related games."""
    article = "https://www.ign.com/articles/2018/07/25/mega-man-x-legacy-collection-2-review"
    if review_url != article:
        return primary, None
    correct_id = "d74cefeb-0a7e-47a5-bd35-20f91a608369"
    correct_url = "https://www.ign.com/games/mega-man-x-legacy-collection-2"
    correct_name = "Mega Man X Legacy Collection 2"
    primary_url = ign_url(primary.get("url") or "", "games")
    # If IGN repairs the source itself, the normal primary-object path applies.
    if (primary.get("id") == correct_id and primary_url == correct_url
            and primary.get("metadata", {}).get("names", {}).get("name") == correct_name):
        return primary, None
    expected_id = "86b6186e-3139-452f-9b63-afc8e352a3f4"
    expected_url = "https://www.ign.com/games/mega-man-x-legacy-collection-1-2"
    if (primary.get("id") != expected_id or primary_url != expected_url
            or page.get("title") != correct_name + " Review - IGN"
            or as_date(page.get("publishDate")) != "2018-07-25"
            or (page.get("review") or {}).get("score") != 6
            or not (page.get("review") or {}).get("verdict", "").startswith(correct_name)):
        raise ValueError("Reviewed article identity has changed; manual check required.")
    candidates = [o for o in page.get("objects") or []
                  if isinstance(o, dict) and o.get("id") == correct_id]
    if len(candidates) != 1:
        raise ValueError("Verified Collection 2 object missing or duplicated; no correction applied.")
    resolved = candidates[0]
    if (resolved.get("type") != "Game"
            or ign_url(resolved.get("url") or "", "games") != correct_url
            or resolved.get("metadata", {}).get("names", {}).get("name") != correct_name):
        raise ValueError("Verified Collection 2 object has changed; manual check required.")
    return resolved, {
        "rule": "mega-man-x-legacy-collection-2-article-v1",
        "review_url": review_url,
        "source": "Manually verified article title, verdict, body (X5-X8) and page.objects",
        "from_game_id": primary["id"], "from_game_url": primary_url,
        "to_game_id": correct_id, "to_game_url": correct_url,
    }


def parse_review(html: str) -> dict:
    parser = ScriptParser()
    parser.feed(html)
    payload = None
    for attrs, text in parser.scripts:
        if attrs.get("id") == "__NEXT_DATA__":
            payload = json.loads(text)
            break
    if payload is None:
        raise ValueError("No __NEXT_DATA__ found. Save the actual review, not a blocked page.")
    page = payload.get("props", {}).get("pageProps", {}).get("page", {})
    primary = page.get("primaryObject") or {}
    obj = primary
    review = page.get("review") or {}
    # Editorial verticals ("Games", "Legacy", "Entertainment") do not
    # determine the reviewed object type. Validate only the primary Game;
    # identity, game URL, score and matching schema are checked below.
    if obj.get("type") != "Game":
        raise ValueError("This page is not a supported game review.")
    name = obj.get("metadata", {}).get("names", {}).get("name")
    if not name or not obj.get("id") or review.get("score") is None:
        raise ValueError("Missing primary game identity or IGN score. No record was inserted.")
    raw_score = review["score"]
    if isinstance(raw_score, bool):
        raise ValueError("Invalid IGN score.")
    score = float(raw_score)
    if not math.isfinite(score) or not 0 <= score <= 10:
        raise ValueError("IGN score is outside the expected 0–10 range.")
    review_url = ign_url(page.get("canonical") or page.get("url") or "", "articles")
    obj, correction = resolve_review_object(page, primary, review_url)
    name = obj.get("metadata", {}).get("names", {}).get("name")
    game_url = ign_url(obj.get("url") or "", "games")
    # Preserve the original primary-object score check even for the exception.
    schema_game_urls = {game_url, ign_url(primary.get("url") or "", "games")}
    for schema in page.get("schema") or []:
        if not isinstance(schema, dict) or schema.get("@type") != "Review":
            continue
        item = schema.get("itemReviewed") or {}
        if item.get("url") and ign_url(item["url"], "games") in schema_game_urls:
            rating = (schema.get("reviewRating") or {}).get("ratingValue")
            if rating is not None and float(rating) != score:
                raise ValueError("Page score and structured-data score disagree; manual check required.")
    releases = []
    for region in obj.get("objectRegions") or []:
        for release in region.get("releases") or []:
            platforms = [p.get("name") or p.get("slug") for p in release.get("platformAttributes") or []]
            for platform in unique(platforms) or [None]:
                releases.append({
                    "region": region.get("region"), "platform": platform,
                    "date": as_date(release.get("date")),
                    "raw_date": release.get("date"),
                    "estimated": bool(release.get("estimatedDate")),
                    "timeframe_year": release.get("timeframeYear"),
                })
    # These are article tags, not a claim of complete worldwide platform coverage.
    attrs = page.get("attributes") or []
    platforms = unique([r["platform"] for r in releases] + [
        (a.get("attribute") or {}).get("name") for a in attrs if a.get("type") == "Platform"])
    genres = unique([(a.get("attribute") or {}).get("name") for a in attrs if a.get("type") == "Genre"])
    if not genres:
        genres = unique([g.get("name") or g.get("slug") for g in obj.get("genres") or []])
    confirmed_us = [r["date"] for r in releases if r["region"] == "US" and r["date"] and not r["estimated"]]
    earliest = min(confirmed_us) if confirmed_us else None
    warnings = []
    raw = {"primaryObject": primary, "review": review, "attributes": attrs}
    if correction:
        warnings.append("Reviewed article identity correction: Collection 1 & 2 -> Collection 2")
        raw["resolvedObject"] = obj
        raw["identityCorrection"] = correction
    for field, value in [("platforms", platforms), ("genres", genres), ("confirmed US release date", earliest)]:
        if not value:
            warnings.append("Missing " + field)
    return {
        "game_id": str(obj["id"]), "name": name, "game_url": game_url,
        "review_url": review_url, "score": score,
        "review_date": as_date(page.get("publishDate")),
        "reviewed_on": review.get("reviewedOn"), "platforms": platforms,
        "genres": genres, "us_release_date": earliest, "releases": releases,
        "warnings": warnings, "captured_at": datetime.now(timezone.utc).isoformat(),
        "raw": raw,
    }


def save_record(connection, record):
    connection.executescript("""
        PRAGMA foreign_keys = ON;
        CREATE TABLE IF NOT EXISTS games (
          game_id TEXT PRIMARY KEY, name TEXT NOT NULL, game_url TEXT NOT NULL,
          platforms_json TEXT NOT NULL, genres_json TEXT NOT NULL, us_release_date TEXT);
        CREATE TABLE IF NOT EXISTS reviews (
          review_url TEXT PRIMARY KEY, game_id TEXT NOT NULL REFERENCES games(game_id),
          score REAL NOT NULL CHECK(score BETWEEN 0 AND 10), review_date TEXT,
          reviewed_on TEXT, captured_at TEXT NOT NULL, warnings_json TEXT NOT NULL,
          raw_json TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS releases (
          game_id TEXT NOT NULL REFERENCES games(game_id), region TEXT, platform TEXT,
          release_date TEXT, raw_date TEXT, estimated INTEGER NOT NULL, timeframe_year TEXT);
    """)
    dump = lambda value: json.dumps(value, ensure_ascii=False)
    with connection:
        connection.execute("""INSERT INTO games VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(game_id) DO UPDATE SET name=excluded.name, game_url=excluded.game_url,
            platforms_json=excluded.platforms_json, genres_json=excluded.genres_json,
            us_release_date=excluded.us_release_date""", (
            record["game_id"], record["name"], record["game_url"], dump(record["platforms"]),
            dump(record["genres"]), record["us_release_date"]))
        connection.execute("""INSERT INTO reviews VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(review_url) DO UPDATE SET game_id=excluded.game_id,
            score=excluded.score, review_date=excluded.review_date, reviewed_on=excluded.reviewed_on,
            captured_at=excluded.captured_at, warnings_json=excluded.warnings_json, raw_json=excluded.raw_json""", (
            record["review_url"], record["game_id"], record["score"], record["review_date"],
            record["reviewed_on"], record["captured_at"], dump(record["warnings"]), dump(record["raw"])))
        connection.execute("DELETE FROM releases WHERE game_id=?", (record["game_id"],))
        connection.executemany("INSERT INTO releases VALUES (?, ?, ?, ?, ?, ?, ?)", [
            (record["game_id"], r["region"], r["platform"], r["date"], r["raw_date"],
             int(r["estimated"]), r["timeframe_year"]) for r in record["releases"]])


HEADERS = ["游戏名称", "IGN评分", "平台（页面所列）", "美区最早发售日期（非预计）", "类型",
           "评测日期", "评测平台", "评测链接", "游戏链接", "采集时间（UTC）", "缺失字段"]


def select_rows(connection, year=None):
    query = """SELECT g.name, r.score, g.platforms_json, g.us_release_date, g.genres_json,
        r.review_date, r.reviewed_on, r.review_url, g.game_url, r.captured_at, r.warnings_json
        FROM reviews r JOIN games g USING(game_id)"""
    params = ()
    if year is not None:
        query += " WHERE substr(g.us_release_date, 1, 4) = ?"
        params = (str(year),)
    query += " ORDER BY g.us_release_date, g.name, r.review_date"
    rows = []
    for source in connection.execute(query, params):
        row = list(source)
        for i in (2, 4, 10):
            row[i] = "; ".join(json.loads(row[i]))
        rows.append(row)
    return rows


def safe_csv(value):
    # Prevent scraped strings being interpreted as Excel formulas.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_files(connection, output, year, csv_only=False):
    output.mkdir(parents=True, exist_ok=True)
    name = f"ign_{year}" if year else "ign_all"
    rows = select_rows(connection, year)
    csv_path = output / (name + ".csv")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)
        writer.writerows([[safe_csv(v) for v in row] for row in rows])
    print(f"CSV: {csv_path} ({len(rows)} reviews)")
    if csv_only:
        return
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ValueError("CSV saved. Install openpyxl to export .xlsx: python -m pip install openpyxl") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IGN评测"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
        for i, value in enumerate(row, 1):
            cell = sheet.cell(sheet.max_row, i)
            if isinstance(value, str):
                cell.data_type = "s"
            if i in (4, 6) and value:
                cell.value = date.fromisoformat(value)
                cell.number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="203864")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    for column, width in zip("ABCDEFGHIJK", [36, 12, 32, 24, 26, 16, 20, 56, 52, 32, 34]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    path = output / (name + ".xlsx")
    workbook.save(path)
    print("Excel:", path)


def fetch_interactive(url):
    from playwright.sync_api import Error, sync_playwright
    url = ign_url(url, "articles")
    os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", r"F:\PlaywrightBrowsers")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False)
        try:
            page = browser.new_page()
            response = page.goto(url, wait_until="domcontentloaded", timeout=60000)
            if response and response.status >= 400:
                raise ValueError(f"IGN returned HTTP {response.status}; stopped.")
            print("Check article and final score in the browser. Scroll if needed.")
            if input("Article and score visible? [y/n]: ").strip().lower() != "y":
                raise ValueError("Stopped without saving a review.")
            ign_url(page.url, "articles")
            return page.content()
        except Error as exc:
            raise ValueError(f"Browser operation failed: {exc}") from exc
        finally:
            browser.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group()
    inputs.add_argument("--html", type=Path, help="Import a saved review.html")
    inputs.add_argument("--latest-sample", action="store_true", help="Use latest data/samples/**/review.html")
    inputs.add_argument("--url", help="Open one review in Chromium")
    inputs.add_argument("--export-only", action="store_true", help="Export existing SQLite data")
    parser.add_argument("--year", type=int, help="Export by earliest confirmed US release year; not review year")
    parser.add_argument("--csv-only", action="store_true", help="Skip .xlsx export")
    parser.add_argument("--data-dir", type=Path, default=ROOT / "data")
    args = parser.parse_args()
    if args.year is not None and not 1900 <= args.year <= 2200:
        parser.error("--year must be between 1900 and 2200")
    args.data_dir.mkdir(parents=True, exist_ok=True)
    database = args.data_dir / "games.sqlite3"
    try:
        if args.export_only and not database.exists():
            raise ValueError("No database yet. Import or collect a review first.")
        record = None
        if not args.export_only:
            path = args.html
            if args.latest_sample:
                candidates = list((args.data_dir / "samples").glob("**/review.html"))
                if not candidates:
                    raise ValueError("No samples found. Copy review.html beside the script and use --html review.html.")
                path = max(candidates, key=lambda p: p.stat().st_mtime_ns)
            if path:
                print("Reading:", path)
                html = path.read_text(encoding="utf-8-sig")
            else:
                html = fetch_interactive(args.url or input("Paste IGN review URL: ").strip())
            record = parse_review(html)
            if not path:
                samples = args.data_dir / "samples" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                samples.mkdir(parents=True)
                (samples / "review.html").write_text(html, encoding="utf-8")
            print("Parsed:", record["name"], "| IGN", record["score"], "| US", record["us_release_date"])
            if record["warnings"]:
                print("Missing fields:", "; ".join(record["warnings"]))
        with sqlite3.connect(database) as connection:
            if record:
                save_record(connection, record)
                print("SQLite saved:", database)
            export_files(connection, args.data_dir / "exports", args.year, args.csv_only)
        return 0
    except PermissionError as exc:
        print("ERROR:", exc)
        print("Close Excel if it has the export open, then retry with --export-only.")
        return 1
    except (ValueError, KeyError, TypeError, OSError, sqlite3.Error) as exc:
        print("ERROR:", exc)
        print("Send this error text if it persists.")
        return 1
    except ImportError as exc:
        print("Missing dependency:", exc)
        print("Install playwright and openpyxl in the project's .venv.")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
